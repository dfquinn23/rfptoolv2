# ui/pages/process_rfp.py
# Page 1: Upload a new RFP, run the full pipeline, download results.

import streamlit as st
import os
from tempfile import NamedTemporaryFile
from docx import Document
from docx.shared import Pt, RGBColor
from pipeline.question_detector import detect_questions
from pipeline.matcher import find_best_match
from pipeline.router import route_question, TIER_COLORS, TIER_MANUAL, TIER_AI_GENERATED
from pipeline.doc_rag import doc_rag_answer
from core.config import OUTPUT_DIR
from core.logger import log_match, log_error


def render():
    st.header("📥 Process New RFP")
    st.write("Upload a raw RFP document. The tool will detect questions, find the best matching past answers, and generate a draft.")

    uploaded = st.file_uploader("Upload .docx RFP", type="docx", key="new_rfp")
    if not uploaded:
        return

    if st.button("Run Pipeline", type="primary"):
        with NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(uploaded.getbuffer())
            tmp_path = tmp.name

        try:
            _run_and_display(tmp_path, uploaded.name)
        finally:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)


def _run_and_display(tmp_path: str, original_name: str):
    results = []

    with st.spinner("Detecting questions..."):
        questions = detect_questions(tmp_path)

    if not questions:
        st.error("No questions detected in this document.")
        return

    st.info(f"Found **{len(questions)} questions**. Matching against knowledge base...")

    progress = st.progress(0)
    status = st.empty()

    for i, q_obj in enumerate(questions):
        question = q_obj["question"]
        status.text(f"Processing {i+1}/{len(questions)}: {question[:80]}...")

        match = find_best_match(question)
        result = route_question(question, match, doc_fallback_fn=doc_rag_answer)
        results.append(result)

        try:
            if result["tier"] != TIER_MANUAL:
                log_match(question, result["answer"], result["final_score"], result["source"], result["tier"])
        except Exception:
            pass

        progress.progress((i + 1) / len(questions))

    status.text("✅ Pipeline complete")
    st.success(f"Processed {len(questions)} questions.")

    # Display results inline
    st.markdown("---")
    for i, r in enumerate(results, 1):
        color = TIER_COLORS.get(r["tier"], "gray")
        st.markdown(f"**Q{i}: {r['question']}**")
        st.markdown(f":{color}[{r['label']}] — Score: {r['final_score']:.2f} | Source: {r['source']}")
        if r["answer"]:
            st.write(r["answer"])
        else:
            st.write("_No answer found — manual response required._")
        st.markdown("---")

    # Generate downloadable .docx
    _generate_output_docx(results, original_name)


def _generate_output_docx(results: list, original_name: str):
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    full_doc = Document()
    full_doc.add_heading("RFP Draft Responses", level=1)

    review_doc = Document()
    review_doc.add_heading("⚠️ Needs Review / Manual Required", level=1)

    for i, r in enumerate(results, 1):
        # Full doc
        p_q = full_doc.add_paragraph()
        run = p_q.add_run(f"Q{i}: {r['question']}")
        run.bold = True
        run.font.size = Pt(11)

        p_label = full_doc.add_paragraph()
        p_label.add_run(f"[{r['label']} | Score: {r['final_score']:.2f} | Source: {r['source']}]").italic = True

        full_doc.add_paragraph(r["answer"] or "— No answer found —")
        full_doc.add_paragraph()

        # Review doc — include anything not auto-verified
        if r["tier"] != "verified":
            rp_q = review_doc.add_paragraph()
            rp_q.add_run(f"Q{i}: {r['question']}").bold = True
            rp_label = review_doc.add_paragraph()
            rp_label.add_run(f"[{r['label']}]").italic = True
            review_doc.add_paragraph(r["answer"] or "— No answer found —")
            review_doc.add_paragraph()

    stem = original_name.replace(".docx", "")
    full_path = os.path.join(OUTPUT_DIR, f"{stem}_draft.docx")
    review_path = os.path.join(OUTPUT_DIR, f"{stem}_needs_review.docx")

    full_doc.save(full_path)
    review_doc.save(review_path)

    col1, col2 = st.columns(2)
    with col1:
        with open(full_path, "rb") as f:
            st.download_button("⬇️ Download Full Draft", f,
                               file_name=f"{stem}_draft.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    with col2:
        with open(review_path, "rb") as f:
            st.download_button("⚠️ Download Review Draft", f,
                               file_name=f"{stem}_needs_review.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
