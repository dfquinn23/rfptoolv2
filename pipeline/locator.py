"""
pipeline/locator.py

Question/Answer-Slot Locator — first stage of the merge/export pipeline.

Given the ORIGINAL incoming RFP document (the one the client sent, before
any editing), this agent identifies, for every question in the document,
exactly where its answer belongs.

Deliberately structure-agnostic. RFP formatting varies by client and cannot
be reliably hardcoded (numbered prefixes, unlabeled questions grouped under
a section header, inline brackets, tables, grouped sub-questions with no
individual answer slot at all, etc. — all occur in practice). Rather than
encode format-specific rules, this agent walks the document's paragraphs in
order, preserving PARAGRAPH INDEX as the one anchor that is never ambiguous,
and asks an LLM to reason about local structure the way a human reader would.

CRITICAL ARCHITECTURAL BOUNDARY — same principle as the rest of the tool:
This agent's output is a LOCATION MAP ONLY. It never touches, rewrites, or
generates answer text. It identifies WHERE an answer goes; a separate,
deterministic Answer Inserter (next pipeline stage) writes the verbatim
approved answer text into that location using python-docx. The LLM here
is scoped to structural understanding, not content generation.

VALIDATION GUARD:
The LLM's structural claims are NOT trusted blindly. After the LLM proposes
a location for each question, a deterministic check (no second LLM call)
verifies the claim against the actual document content:
  - "blank_line" claims are only accepted if the target paragraph really is
    blank in the source document.
  - "inline_bracket" claims are only accepted if the answer paragraph is the
    same paragraph as the question.
  - Two questions claiming the same answer slot are both rejected.
Anything that fails validation is downgraded to "unclear" rather than acted
on. This exists because the LLM will confidently propose a plausible-looking
but WRONG location (e.g. pointing one question's answer at the next
question's own paragraph, inside a tightly-packed cluster of sub-questions
with no individual blank line) — and a wrong location, if trusted, means the
Answer Inserter silently overwrites another question with someone else's
answer. "Unclear" is always safer than a wrong guess; unclear/failed
questions fall through to manual review / the addendum, never into a wrong
paragraph.

Tables are intentionally NOT processed by this agent — python-docx does not
index table-cell paragraphs the same way as body paragraphs, and per project
decision, table/chart fields are skipped for manual completion regardless.

TABLE-ADJACENCY CHECK:
A blank paragraph that sits directly above a real .docx table (e.g. "Complete
the table below...", then a blank line, then the actual table) will look
identical to a normal dedicated answer paragraph from the LLM's point of
view — both are just "a blank paragraph right after the question." Left
unchecked, this passes the blank_line validation above even though the true
answer destination is the table, not that blank line. This agent walks the
document body in order (not just doc.paragraphs) to know exactly which
paragraph indices sit immediately before a table, and relabels any claim
that lands on one of those paragraphs as "table" instead of "blank_line" —
routing it to manual completion instead of a false-positive auto-insert.

Public API:
    locate_answer_slots(file_path, verbose=False) -> list[dict]
        Raw, validated location records for every question the agent found
        in the document. See _RECORD SCHEMA below.

    get_structure_verification(file_path, approved_questions=None, verbose=False) -> list[dict]
        Built for the "Verify RFP Structure" UI step. Same underlying
        location work, but shaped for review: one row per question (matched
        against the user's approved/edited question list when provided),
        each tagged "verified" or "needs_review", with a short surrounding-
        paragraph snippet for context so a human can make a fast call.

_RECORD SCHEMA (locate_answer_slots):
    {
        "question_text":      str,
        "question_para_idx":  int,
        "answer_para_idx":    int | None,
        "insertion_type":     "blank_line" | "inline_bracket" | "table" | "unclear",
    }
    "table" means the answer belongs in a Word table, not a paragraph — this
    agent does not write locations inside tables; it always routes these to
    manual completion (answer_para_idx is always None for "table").
"""

from __future__ import annotations

import difflib
import json
import os
import re
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from openai import OpenAI

load_dotenv()

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

OPENAI_MODEL     = "gpt-4o-mini"
MAX_CHUNK_LINES  = 80    # split documents larger than this many indexed lines.
                         # Kept intentionally small: this bounds not just the
                         # LLM's input but, more importantly, its OUTPUT — a
                         # dense document (many short questions) can need a
                         # JSON array far larger than a sparse one with the
                         # same paragraph count. A chunk that's too large can
                         # cause the response to be truncated mid-generation,
                         # which silently discards every question in that
                         # chunk (see _parse_json_array_lenient below for the
                         # safety net on top of this).
CHUNK_OVERLAP    = 10    # lines of overlap between chunks, to avoid cutting
                         # a question away from its answer slot at a boundary
MAX_OUTPUT_TOKENS = 8192 # generous ceiling for the JSON response itself
CONTEXT_WINDOW   = 2     # paragraphs of context shown before/after a question
                         # in the structure-verification snippet
MATCH_THRESHOLD  = 0.6   # min similarity to match an approved question back
                         # to a located record (difflib ratio, 0-1)


# ---------------------------------------------------------------------------
# Document structure extraction
# ---------------------------------------------------------------------------

def _extract_document_structure(
    file_path: str,
) -> tuple[list[tuple[int, str]], set[int], int]:
    """
    Walk the document once, returning everything the Locator needs:

      1. Indexed paragraphs — [(idx, text), ...] over doc.paragraphs, in the
         exact index the Answer Inserter will later use to write into this
         same document. Blank paragraphs ARE included and ARE meaningful —
         a blank paragraph is very often the answer slot itself.

      2. table_adjacent_idxs — the set of paragraph indices that sit
         DIRECTLY above a real Word table in the document body (no other
         paragraph in between). See module docstring, "TABLE-ADJACENCY
         CHECK" — these paragraphs look like ordinary blank answer slots
         but are actually just the gap before a table.

      3. table_count — total number of tables in the document (FYI/logging
         only).

    Table-cell paragraphs are NOT included in the indexed paragraph list —
    see module docstring.
    """
    doc = Document(file_path)
    indexed = [(i, para.text.strip()) for i, para in enumerate(doc.paragraphs)]

    table_adjacent_idxs: set[int] = set()
    para_idx = -1
    last_child_was_paragraph = False

    for child in doc.element.body.iterchildren():
        tag = child.tag
        if tag.endswith("}p"):
            para_idx += 1
            last_child_was_paragraph = True
        elif tag.endswith("}tbl"):
            if last_child_was_paragraph:
                table_adjacent_idxs.add(para_idx)
            last_child_was_paragraph = False
        else:
            # sectPr and other structural elements — neither a paragraph
            # nor a table; doesn't carry the "directly above a table" flag
            # forward.
            last_child_was_paragraph = False

    return indexed, table_adjacent_idxs, len(doc.tables)


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------

def _chunk_indexed_paragraphs(
    indexed: list[tuple[int, str]],
    max_lines: int = MAX_CHUNK_LINES,
    overlap: int = CHUNK_OVERLAP,
) -> list[list[tuple[int, str]]]:
    """Split the indexed paragraph list into overlapping chunks."""
    if len(indexed) <= max_lines:
        return [indexed]

    chunks = []
    start = 0
    while start < len(indexed):
        end = min(start + max_lines, len(indexed))
        chunks.append(indexed[start:end])
        if end == len(indexed):
            break
        start = end - overlap
    return chunks


# ---------------------------------------------------------------------------
# LLM call
# ---------------------------------------------------------------------------

_SYSTEM_PROMPT = """\
You are a document analyst specialising in RFP (Request for Proposal) documents.

You will be given a numbered list of paragraphs from an incoming RFP, in the
format "IDX: text". Some paragraphs are blank (empty text after the colon) —
this is meaningful, not an error; blank paragraphs are frequently the space
left for an answer.

Your task: identify every question a respondent must answer, and for each
one, determine exactly which paragraph its answer belongs in.

Questions may be formatted in ANY of these ways — do not assume one style
applies to the whole document, and do not rely on numbering patterns being
consistent:
  - Individually numbered ("1.1.1 What is your firm's AUM?")
  - Grouped under a section header with NO individual numbering, each still
    getting its own separate answer slot
  - Plain prose ending in "?"
  - Imperative phrasing ("Please describe...", "Provide an overview of...")
  - Inline blank brackets on the same line ("Founded year: [ ]")
  - Tightly packed clusters of sub-questions with NO blank line between any
    of them (e.g. a header question followed by several one-line sub-items,
    with only a single blank line after the whole cluster, not one per item)

For each question, determine the answer location using LOCAL structure:
  - "blank_line": the answer belongs in a separate (usually blank) paragraph
    that immediately follows the question, before the next question begins.
    Only use this when you are confident that specific paragraph is truly
    dedicated to this question's answer alone.
  - "inline_bracket": the answer belongs inline on the same paragraph as the
    question (e.g. after a colon or in brackets on that same line).
  - "unclear": you cannot confidently determine a DEDICATED answer location
    for this specific question. This includes cases where several
    sub-questions are packed together with no individual blank line — do
    NOT guess by pointing one question's answer at the next question's own
    paragraph. A wrong location is worse than an honest "unclear".

Do not confuse one question's answer slot with the paragraph(s) belonging to
the NEXT question. If the paragraph immediately following a question is
itself clearly another question (not blank, not an answer), you MUST mark
that question "unclear" rather than pointing at the next question's text.

Ignore section headers on their own (e.g. "FIRM HISTORY" is not a question),
instructions, deadlines, cover-page text, and legal disclaimers — these are
not questions and should not appear in your output.

Return ONLY a valid JSON array of objects, no preamble, no markdown fences:
[
  {
    "question_text": "exact question text as it appears",
    "question_para_idx": <int>,
    "answer_para_idx": <int or null>,
    "insertion_type": "blank_line" | "inline_bracket" | "unclear"
  }
]

If insertion_type is "unclear", answer_para_idx should be null.
"""


def _parse_json_array_lenient(raw: str) -> list[dict]:
    """
    Parse the LLM's JSON array response, tolerating truncation.

    If the response got cut off mid-generation (hit the token ceiling before
    finishing), a strict json.loads() fails and — without this — the entire
    response gets thrown away, including every complete, valid question
    object that came before the cutoff. That's a silent, total data loss for
    that whole chunk, not a graceful degradation.

    Instead: scan for complete top-level {...} objects by bracket depth, and
    keep every one that parses cleanly, discarding only the incomplete tail.
    A shrunk chunk size and a generous MAX_OUTPUT_TOKENS (see config above)
    should make truncation rare — this is the defense-in-depth backstop for
    when it happens anyway.
    """
    raw = raw.strip()
    raw = re.sub(r"^```[a-z]*\n?", "", raw)
    raw = re.sub(r"\n?```$",        "", raw)

    try:
        parsed = json.loads(raw)
        if isinstance(parsed, list):
            return parsed
    except json.JSONDecodeError:
        pass

    objects: list[dict] = []
    depth = 0
    start = None
    for i, ch in enumerate(raw):
        if ch == "{":
            if depth == 0:
                start = i
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0 and start is not None:
                candidate = raw[start:i + 1]
                try:
                    objects.append(json.loads(candidate))
                except json.JSONDecodeError:
                    pass  # this one object was itself malformed — skip it only
                start = None

    if objects:
        print(f"[locator] WARNING: response was truncated or malformed — "
              f"recovered {len(objects)} complete question(s) from it rather "
              f"than discarding the whole chunk. Consider this a signal to "
              f"check MAX_CHUNK_LINES if it happens often.")
    else:
        print(f"[locator] WARNING: LLM returned non-JSON, nothing recoverable:\n{raw[:200]}")

    return objects


def _call_llm(client: OpenAI, indexed_chunk: list[tuple[int, str]]) -> list[dict]:
    """Send one indexed chunk to the LLM and parse the returned JSON array."""
    text_block = "\n".join(f"{idx}: {text}" for idx, text in indexed_chunk)

    response = client.chat.completions.create(
        model=OPENAI_MODEL,
        temperature=0,
        max_tokens=MAX_OUTPUT_TOKENS,
        messages=[
            {"role": "system", "content": _SYSTEM_PROMPT},
            {"role": "user",   "content": text_block},
        ],
    )

    raw = response.choices[0].message.content.strip()
    return _parse_json_array_lenient(raw)


# ---------------------------------------------------------------------------
# Validation guard — deterministic, no LLM involved
# ---------------------------------------------------------------------------

def _validate_records(
    records: list[dict],
    indexed: list[tuple[int, str]],
    table_adjacent_idxs: set[int] = frozenset(),
) -> list[dict]:
    """
    Verify each LLM-proposed location against the actual document content.
    Anything that doesn't check out is downgraded rather than trusted.
    See module docstring for rationale.
    """
    text_by_idx = dict(indexed)
    validated: list[dict] = []

    def _as(rec: dict, new_type: str, new_aidx=None) -> dict:
        return {**rec, "insertion_type": new_type, "answer_para_idx": new_aidx}

    for r in records:
        qidx  = r.get("question_para_idx")
        aidx  = r.get("answer_para_idx")
        itype = r.get("insertion_type")

        if qidx is None or qidx not in text_by_idx:
            validated.append(_as(r, "unclear"))
            continue

        if itype == "blank_line":
            if aidx is None or aidx not in text_by_idx or text_by_idx[aidx] != "":
                validated.append(_as(r, "unclear"))
                continue
            if aidx in table_adjacent_idxs:
                # This "blank line" is really just the gap before a table —
                # the true answer destination is the table, not this
                # paragraph. Route to manual completion, distinctly labeled.
                validated.append(_as(r, "table"))
                continue

        elif itype == "inline_bracket":
            if aidx != qidx:
                validated.append(_as(r, "unclear"))
                continue

        elif itype == "unclear":
            # Give a more specific, more useful label when we can tell WHY
            # it's unclear — the question sits directly above a table.
            if qidx in table_adjacent_idxs:
                validated.append(_as(r, "table"))
            else:
                validated.append({**r, "answer_para_idx": None})
            continue

        else:
            # Unrecognized insertion_type from the LLM — treat as unverified
            validated.append(_as(r, "unclear"))
            continue

        validated.append(r)

    # Collision check: two questions claiming the same answer paragraph.
    # Neither can be trusted — both get downgraded.
    slot_claims: dict[int, list[int]] = {}
    for i, r in enumerate(validated):
        if r["insertion_type"] in ("blank_line", "inline_bracket") and r.get("answer_para_idx") is not None:
            slot_claims.setdefault(r["answer_para_idx"], []).append(i)

    for idx, claimant_positions in slot_claims.items():
        if len(claimant_positions) > 1:
            for pos in claimant_positions:
                validated[pos] = {
                    **validated[pos],
                    "insertion_type": "unclear",
                    "answer_para_idx": None,
                }

    return validated


# ---------------------------------------------------------------------------
# Core location logic (shared by both public entry points)
# ---------------------------------------------------------------------------

def _locate_from_indexed(
    indexed: list[tuple[int, str]],
    table_adjacent_idxs: set[int] = frozenset(),
    verbose: bool = False,
) -> list[dict]:
    if not indexed:
        return []

    chunks = _chunk_indexed_paragraphs(indexed)
    client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
    all_records: list[dict] = []

    for i, chunk in enumerate(chunks, 1):
        if verbose or len(chunks) > 1:
            print(f"[locator] Processing chunk {i}/{len(chunks)} "
                  f"({len(chunk)} lines, idx {chunk[0][0]}-{chunk[-1][0]})...")
        records = _call_llm(client, chunk)
        all_records.extend(records)
        if verbose:
            print(f"  → {len(records)} question(s) located in this chunk.")

    # Deduplicate — overlapping chunks can surface the same question twice.
    # Keep the first occurrence (earlier chunk = more context before it).
    seen: set[tuple[str, int]] = set()
    unique: list[dict] = []
    for r in all_records:
        key = (str(r.get("question_text", "")).lower().strip(),
               r.get("question_para_idx"))
        if key not in seen:
            seen.add(key)
            unique.append(r)

    validated = _validate_records(unique, indexed, table_adjacent_idxs)

    table_count   = sum(1 for r in validated if r.get("insertion_type") == "table")
    unclear_count = sum(1 for r in validated if r.get("insertion_type") == "unclear")
    print(f"[locator] Done. {len(validated)} question(s) located "
          f"({unclear_count} unclear, {table_count} route to table/manual "
          f"after validation).")

    return validated


# ---------------------------------------------------------------------------
# Public API — raw location map
# ---------------------------------------------------------------------------

def locate_answer_slots(file_path: str, verbose: bool = False) -> list[dict]:
    """
    Identify where each question's answer belongs in the original RFP.

    Args:
        file_path: Path to the ORIGINAL incoming RFP .docx file (the same
                    document the Answer Inserter will later write into —
                    paragraph indices must match).
        verbose:    Print chunk-level progress.

    Returns:
        List of validated location records (see module docstring schema).
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"[locator] File not found: {file_path}")
    if path.suffix.lower() != ".docx":
        raise ValueError(f"[locator] Expected .docx file, got: {path.suffix}")

    print(f"[locator] Reading: {path.name}")
    indexed, table_adjacent_idxs, table_count = _extract_document_structure(str(path))

    if not indexed:
        print("[locator] WARNING: No paragraphs found in document.")
        return []

    if table_count:
        print(f"[locator] NOTE: Document contains {table_count} table(s), "
              f"{len(table_adjacent_idxs)} paragraph(s) sit directly above "
              f"one. Those route to manual completion, not auto-insert.")

    print(f"[locator] Indexed {len(indexed)} paragraph(s).")

    return _locate_from_indexed(indexed, table_adjacent_idxs, verbose=verbose)


# ---------------------------------------------------------------------------
# Public API — structure verification (for the "Verify RFP Structure" step)
# ---------------------------------------------------------------------------

def _context_snippet(indexed: list[tuple[int, str]], center_idx: int,
                      window: int = CONTEXT_WINDOW) -> str:
    """Build a small surrounding-paragraph snippet for display in the UI."""
    lo = max(0, center_idx - window)
    hi = min(len(indexed) - 1, center_idx + window)
    lines = []
    for idx, text in indexed[lo:hi + 1]:
        marker  = ">> " if idx == center_idx else "   "
        display = text if text else "(blank)"
        lines.append(f"{marker}{idx}: {display}")
    return "\n".join(lines)


def match_to_approved(records: list[dict], approved_questions: list[str],
                       threshold: float = MATCH_THRESHOLD) -> list[dict]:
    """
    Align located records to a list of question strings (fuzzy text match).
    Public — shared by get_structure_verification() (below) and by
    pipeline.inserter, so the question-matching logic lives in exactly one
    place rather than being duplicated across the Locator and the Inserter.

    Questions with no confident match get a placeholder "unclear" record —
    they still need to show up downstream, since they'll fall to the
    addendum either way.
    """
    used: set[int] = set()
    matched: list[dict] = []

    for aq in approved_questions:
        best_idx, best_score = None, 0.0
        for i, r in enumerate(records):
            if i in used:
                continue
            score = difflib.SequenceMatcher(
                None, aq.lower().strip(), str(r.get("question_text", "")).lower().strip()
            ).ratio()
            if score > best_score:
                best_score, best_idx = score, i

        if best_idx is not None and best_score >= threshold:
            used.add(best_idx)
            rec = dict(records[best_idx])
            rec["question_text"] = aq  # display the user's approved wording
            matched.append(rec)
        else:
            matched.append({
                "question_text": aq,
                "question_para_idx": None,
                "answer_para_idx": None,
                "insertion_type": "unclear",
            })

    return matched


def get_structure_verification(
    file_path: str,
    approved_questions: list[str] | None = None,
    verbose: bool = False,
) -> list[dict]:
    """
    Produce the data for the "Verify RFP Structure" UI step.

    Args:
        file_path:           Path to the ORIGINAL incoming RFP .docx file.
        approved_questions:   The user's final, edited question list from
                              Pipeline Stage 2. When provided, records are
                              matched back to this list (fuzzy text match)
                              so the verification UI reflects exactly the
                              questions the user approved — not raw
                              detection noise. When None, all located
                              records are returned as-is.
        verbose:              Print chunk-level progress.

    Returns:
        One dict per question:
            {
                "question_text":     str,
                "question_para_idx": int | None,
                "answer_para_idx":   int | None,
                "insertion_type":    "blank_line" | "inline_bracket" | "table" | "unclear",
                "status":            "verified" | "needs_review",
                "context_snippet":   str,
            }
        "table" is a distinct flavor of needs_review — it means the answer
        belongs in a Word table, so the UI can tell the user WHY, rather
        than just "unclear".
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"[locator] File not found: {file_path}")

    indexed, table_adjacent_idxs, _table_count = _extract_document_structure(str(path))
    records = _locate_from_indexed(indexed, table_adjacent_idxs, verbose=verbose)

    if approved_questions is not None:
        records = match_to_approved(records, approved_questions)

    results: list[dict] = []
    for r in records:
        qidx  = r.get("question_para_idx")
        itype = r.get("insertion_type")
        status = "verified" if itype in ("blank_line", "inline_bracket") else "needs_review"
        snippet = _context_snippet(indexed, qidx) if qidx is not None else "(question not located in document)"

        results.append({
            "question_text":     r.get("question_text", ""),
            "question_para_idx": qidx,
            "answer_para_idx":   r.get("answer_para_idx"),
            "insertion_type":    itype,
            "status":            status,
            "context_snippet":   snippet,
        })

    verified_count = sum(1 for r in results if r["status"] == "verified")
    print(f"[locator] Structure verification: {verified_count}/{len(results)} "
          f"verified, {len(results) - verified_count} need review.")

    return results


# ---------------------------------------------------------------------------
# CLI — smoke test
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

    if len(sys.argv) < 2:
        print("Usage: python -m pipeline.locator <path/to/original_rfp.docx>")
        print("       python -m pipeline.locator <path/to/original_rfp.docx> --verbose")
        print("       python -m pipeline.locator <path/to/original_rfp.docx> --verify")
        sys.exit(1)

    target  = sys.argv[1]
    verbose = "--verbose" in sys.argv or "-v" in sys.argv
    verify  = "--verify" in sys.argv

    if verify:
        results = get_structure_verification(target, verbose=verbose)
        print(f"\n--- Structure Verification ({len(results)}) ---")
        for r in results:
            if r["status"] == "verified":
                icon = "✅"
            elif r["insertion_type"] == "table":
                icon = "🗂️ "
            else:
                icon = "⚠️ "
            qidx = r["question_para_idx"]
            aidx = r["answer_para_idx"]
            print(f"  {icon} [{r['insertion_type']:>14}]  Q@{qidx if qidx is not None else '—':<4} "
                  f"→ A@{aidx if aidx is not None else '—':<4}  {r['question_text'][:65]}")
    else:
        records = locate_answer_slots(target, verbose=verbose)
        print(f"\n--- Located Answer Slots ({len(records)}) ---")
        for r in records:
            qidx  = r.get("question_para_idx")
            aidx  = r.get("answer_para_idx")
            itype = r.get("insertion_type")
            qtext = r.get("question_text", "")
            print(f"  [{itype:>14}]  Q@{qidx:<4} → A@{aidx if aidx is not None else '—':<4}  {qtext[:70]}")