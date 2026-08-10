"""
pipeline/addendum.py

Addendum Builder — fourth stage of the merge pipeline.

Takes everything the Answer Inserter could NOT verify a location for
(InsertResult.skipped) and produces a clean, standalone Word document
listing each such question with its approved answer, grouped by WHY it
needs manual attention. This is the safety net for the whole merge design:
nothing that can't be verified gets guessed at or force-placed into the
original document — it lands here instead, in a form a human can act on
in a couple of minutes.

Adapted from the existing _build_export_doc() in pages/2_Review.py — same
underlying pattern (write approved Q&A pairs to a clean docx), reshaped to
group by insertion category (table / needs manual placement / missing
answer) instead of by confidence tier, since that's the relevant grouping
once the merge step has already run.

No LLM involvement. Deterministic, like the Inserter.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


# ---------------------------------------------------------------------------
# Grouping
# ---------------------------------------------------------------------------

# Order matters — this is the order sections appear in the document.
_GROUP_ORDER = ["missing_answer", "table", "manual_placement"]

_GROUP_LABELS = {
    "missing_answer":   "⚠️ Missing Answer — Needs Attention",
    "table":            "📋 Belongs in a Table",
    "manual_placement": "✏️ Needs Manual Placement",
}

_GROUP_INTROS = {
    "missing_answer": (
        "These questions were approved without answer text attached. "
        "They need an answer before they can go anywhere."
    ),
    "table": (
        "These questions expect an answer inside a table in the original "
        "document, which this tool does not write to automatically. "
        "Please copy each answer into the corresponding table by hand."
    ),
    "manual_placement": (
        "The tool could not confidently determine where these answers "
        "belong in the original document (or couldn't confidently match "
        "them to a question at all). Please place each answer manually."
    ),
}


def _categorize(item: dict) -> str:
    """Sort one skipped item into a display group."""
    if not (item.get("answer_text") or "").strip():
        return "missing_answer"
    if item.get("insertion_type") == "table":
        return "table"
    return "manual_placement"


# ---------------------------------------------------------------------------
# Result type
# ---------------------------------------------------------------------------

@dataclass
class AddendumResult:
    output_path: str
    counts: dict = field(default_factory=dict)  # group -> count

    @property
    def total(self) -> int:
        return sum(self.counts.values())


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def build_addendum(
    skipped_items: list[dict],
    output_path: str,
    source_filename: str = "",
) -> AddendumResult:
    """
    Build the addendum document from everything the Inserter skipped.

    Args:
        skipped_items:    InsertResult.skipped — list of dicts with
                           "question_text", "answer_text", "insertion_type",
                           and "reason".
        output_path:       Where to save the addendum .docx.
        source_filename:   Original RFP filename, shown in the document
                            header for reference.

    Returns:
        AddendumResult with the output path and a per-group item count.
    """
    grouped: dict[str, list[dict]] = {g: [] for g in _GROUP_ORDER}
    for item in skipped_items:
        grouped[_categorize(item)].append(item)

    doc = Document()

    title = doc.add_heading("RFP Addendum — Items Needing Manual Attention", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    meta_lines = []
    if source_filename:
        meta_lines.append(f"Source RFP: {source_filename}")
    meta_lines.append(f"Generated: {datetime.now().strftime('%d %b %Y %H:%M')}")
    total_items = sum(len(v) for v in grouped.values())
    meta_lines.append(f"Total items: {total_items}")

    meta = doc.add_paragraph("    |    ".join(meta_lines))
    meta.runs[0].font.size = Pt(9)
    meta.runs[0].font.italic = True

    doc.add_paragraph(
        "Everything below could not be automatically merged into the "
        "original RFP document. Nothing here was guessed at — each item "
        "is grouped by why it needs a human to finish it."
    )
    doc.add_paragraph("")

    if total_items == 0:
        doc.add_paragraph("Nothing to review — every approved answer was inserted automatically.")

    for group_key in _GROUP_ORDER:
        items = grouped[group_key]
        if not items:
            continue

        doc.add_heading(f"{_GROUP_LABELS[group_key]} ({len(items)})", level=2)
        intro = doc.add_paragraph(_GROUP_INTROS[group_key])
        intro.runs[0].font.italic = True
        intro.runs[0].font.size = Pt(9)
        doc.add_paragraph("")

        for item in items:
            q_para = doc.add_paragraph()
            q_run  = q_para.add_run(item.get("question_text", ""))
            q_run.bold = True
            q_run.font.size = Pt(11)

            answer_text = (item.get("answer_text") or "").strip()
            a_para = doc.add_paragraph(answer_text or "[No answer provided]")
            a_para.paragraph_format.space_after = Pt(4)

            if group_key == "manual_placement":
                # A little breadcrumb for why this one landed here — helps
                # the reviewer decide how much scrutiny it needs.
                reason = item.get("reason", "")
                if reason:
                    r_para = doc.add_paragraph(f"({reason})")
                    r_para.runs[0].font.size    = Pt(8)
                    r_para.runs[0].font.italic  = True

            a_para.paragraph_format.space_after = Pt(14)

        doc.add_paragraph("")

    dest = Path(output_path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))

    counts = {g: len(items) for g, items in grouped.items()}
    print(f"[addendum] Built addendum with {total_items} item(s) "
          f"({', '.join(f'{k}={v}' for k, v in counts.items())}) → {dest}")

    return AddendumResult(output_path=str(dest), counts=counts)


# ---------------------------------------------------------------------------
# CLI — smoke test
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys
    import json

    if len(sys.argv) < 2:
        print("Usage: python -m pipeline.addendum <skipped_items.json> [output.docx]")
        print('  skipped_items.json: InsertResult.skipped, i.e. a list of')
        print('  {"question_text": ..., "answer_text": ..., "insertion_type": ..., "reason": ...}')
        sys.exit(1)

    items_path  = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else "addendum_test_output.docx"

    with open(items_path, "r", encoding="utf-8") as f:
        skipped = json.load(f)

    result = build_addendum(skipped, output_path, source_filename=Path(items_path).stem)

    print(f"\n--- Addendum built: {result.total} item(s) ---")
    for group, count in result.counts.items():
        print(f"  {group}: {count}")
    print(f"\nOutput: {result.output_path}")